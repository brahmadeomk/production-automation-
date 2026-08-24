"""XLSX export of filtered records and summaries (SRS section 14)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from ..db.database import Database
from .engine import Summary

try:  # openpyxl is required only for export, not for programming
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover - exercised on stations without export
    OPENPYXL_AVAILABLE = False


LOG_COLUMNS = (
    ("Timestamp", "Timestamp"),
    ("SerialNo", "Serial No"),
    ("Result", "Result"),
    ("ProjectName", "Project"),
    ("Operator", "Operator"),
    ("FirmwareVersion", "Firmware"),
    ("HwRevision", "HW Rev"),
    ("ProductVariant", "Variant"),
    ("StationId", "Station"),
    ("Signature", "Signature"),
    ("DurationMs", "Cycle (ms)"),
    ("ErrorCode", "Error Code"),
    ("Error", "Error Detail"),
)

_HEADER_FILL = "FF1F3864"
_PASS_FILL = "FFDDF3DD"
_FAIL_FILL = "FFF8D7DA"


class ExcelExportUnavailable(RuntimeError):
    """openpyxl is not installed on this station."""


def _require_openpyxl() -> None:
    if not OPENPYXL_AVAILABLE:
        raise ExcelExportUnavailable(
            "openpyxl is not installed; run: pip install openpyxl"
        )


def _style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, max_width: int = 46) -> None:
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), max_width)


def _write_rows(ws, rows: Sequence[Any]) -> None:
    ws.append([label for _, label in LOG_COLUMNS])
    pass_fill = PatternFill("solid", fgColor=_PASS_FILL)
    fail_fill = PatternFill("solid", fgColor=_FAIL_FILL)
    for row in rows:
        record = dict(row)
        ws.append([record.get(key, "") for key, _ in LOG_COLUMNS])
        fill = pass_fill if record.get("Result") == "PASS" else fail_fill
        ws.cell(row=ws.max_row, column=3).fill = fill
    _style_header(ws)
    _autosize(ws)
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_COLUMNS))}{ws.max_row}"


def export_records(
    rows: Iterable[Any],
    destination: str | Path,
    *,
    title: str = "Production Log",
    filters: Optional[Dict[str, Any]] = None,
) -> Path:
    """Write filtered production records to an XLSX workbook."""
    _require_openpyxl()
    rows = list(rows)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Production Log"
    _write_rows(ws, rows)

    meta = wb.create_sheet("Export Info")
    meta.append(["Field", "Value"])
    meta.append(["Title", title])
    meta.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    meta.append(["Records", len(rows)])
    for key, value in (filters or {}).items():
        if value not in (None, ""):
            meta.append([str(key), str(value)])
    _style_header(meta)
    _autosize(meta)

    wb.save(destination)
    return destination


def export_summary(
    summary: Summary,
    destination: str | Path,
    *,
    rows: Optional[Iterable[Any]] = None,
    company_name: str = "",
) -> Path:
    """Write a period report: totals, per-operator, per-project, failures."""
    _require_openpyxl()
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Report", summary.label])
    if company_name:
        ws.append(["Company", company_name])
    ws.append(["Window start", summary.start])
    ws.append(["Window end", summary.end])
    ws.append([])
    ws.append(["Metric", "Value"])
    header_row = ws.max_row
    for name, value in (
        ("Total programmed", summary.total),
        ("PASS", summary.passed),
        ("FAIL", summary.failed),
        ("Yield %", summary.yield_pct),
        ("Average cycle (s)", round(summary.avg_cycle_ms / 1000.0, 2)),
    ):
        ws.append([name, value])
    _style_header(ws, header_row)
    ws["A1"].font = Font(bold=True, size=14)
    _autosize(ws)

    def add_sheet(name: str, headers: Sequence[str], keys: Sequence[str], data: List[Dict[str, Any]]):
        if not data:
            return
        sheet = wb.create_sheet(name)
        sheet.append(list(headers))
        for item in data:
            sheet.append([item.get(k, "") for k in keys])
        _style_header(sheet)
        _autosize(sheet)

    add_sheet(
        "By Operator",
        ("Operator", "Total", "PASS", "FAIL", "Yield %", "Avg cycle (ms)"),
        ("operator", "total", "passed", "failed", "yield_pct", "avg_cycle_ms"),
        summary.by_operator,
    )
    add_sheet(
        "By Project",
        ("Project", "Total", "PASS", "FAIL", "Yield %", "Avg cycle (ms)"),
        ("project", "total", "passed", "failed", "yield_pct", "avg_cycle_ms"),
        summary.by_project,
    )
    add_sheet(
        "Failures",
        ("Error Code", "Example Message", "Count"),
        ("error_code", "message", "count"),
        summary.by_error,
    )
    add_sheet(
        "By Day",
        ("Day", "Total", "PASS", "FAIL", "Yield %"),
        ("day", "total", "passed", "failed", "yield_pct"),
        summary.by_day,
    )

    if rows is not None:
        detail = wb.create_sheet("Records")
        _write_rows(detail, list(rows))

    wb.save(destination)
    return destination


def default_filename(prefix: str, suffix: str = "") -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    tail = f"-{suffix}" if suffix else ""
    return f"{prefix}{tail}-{stamp}.xlsx"


def export_period(
    db: Database,
    summary: Summary,
    destination: str | Path,
    *,
    include_records: bool = True,
    company_name: str = "",
) -> Path:
    """Convenience wrapper: summary plus the records it covers."""
    rows = (
        db.search_production(date_from=summary.start, date_to=summary.end, limit=1_000_000)
        if include_records
        else None
    )
    return export_summary(summary, destination, rows=rows, company_name=company_name)
